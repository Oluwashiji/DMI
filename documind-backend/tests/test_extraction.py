"""End-to-end tests for the extraction flow.

Real PDFs/images go through the real PDF parser and Tesseract; only the Groq
call is replaced with a canned response. Run with:  pytest -q
(needs: pip install pytest httpx, and the tesseract binary installed)
"""
import io
import json
import os
import tempfile
import uuid

_tmp = tempfile.mkdtemp()
os.environ["DATABASE_URL"] = os.environ.get("TEST_DATABASE_URL", f"sqlite:///{_tmp}/test.db")
os.environ["GROQ_API_KEY"] = "test-key"
os.environ["JWT_SECRET_KEY"] = "test-secret"

import fitz  # noqa: E402
import groq  # noqa: E402
import httpx  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

from app.main import app  # noqa: E402
from app.schemas.extraction import ExtractedInvoice, to_number  # noqa: E402
from app.services import llm, ocr  # noqa: E402

INVOICE_LINES = [
    "ACME SUPPLIES LTD",
    "INVOICE",
    "Invoice No: INV-2026-0042",
    "Date: 2026-09-14",
    "Bill to: Example Customer",
    "Description        Qty   Unit Price   Total",
    "Widget A            2      25.00      50.00",
    "Widget B            1      100.00     100.00",
    "Subtotal: 150.00",
    "VAT 7.5%: 11.25",
    "TOTAL DUE: USD 161.25",
]

GOOD_JSON = json.dumps({
    "vendor_name": "Acme Supplies Ltd",
    "invoice_number": "INV-2026-0042",
    "date": "2026-09-14",
    "currency": "USD",
    "line_items": [
        {"description": "Widget A", "quantity": 2, "unit_price": 25, "total": 50},
        {"description": "Widget B", "quantity": 1, "unit_price": "100.00", "total": "$100.00"},
    ],
    "subtotal": "150.00",
    "tax": 11.25,
    "total": "161.25",
})


# ---------- sample document builders ----------


def text_pdf(lines=INVOICE_LINES, pages=1) -> bytes:
    doc = fitz.open()
    for _ in range(pages):
        page = doc.new_page()
        page.insert_text((50, 72), "\n".join(lines), fontsize=12, fontname="cour")
    return doc.tobytes()


def invoice_png(lines=INVOICE_LINES) -> bytes:
    font = ImageFont.load_default(size=30)
    img = Image.new("RGB", (1400, 80 + 50 * len(lines)), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((40, 40 + 50 * i), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def scanned_pdf() -> bytes:
    """A PDF whose only content is a picture (no text layer), like a scan."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=invoice_png())
    return doc.tobytes()


# ---------- fixtures ----------


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


RUN_ID = uuid.uuid4().hex[:8]  # unique per run, so the suite can be re-run against a database that keeps old accounts


def signup(client, email) -> dict:
    email = email.replace("@", f"+{RUN_ID}@")
    r = client.post("/auth/signup", json={"full_name": "Test User", "email": email, "password": "password123"})
    assert r.status_code == 201, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def headers(client):
    return signup(client, "owner@example.com")


@pytest.fixture
def fake_llm(monkeypatch):
    seen = []

    def _fake(text):
        seen.append(text)
        return GOOD_JSON

    monkeypatch.setattr(llm, "_call_model", _fake)
    return seen


def run_extract(client, headers, name, data, mime="application/octet-stream"):
    """POST to /api/extract and parse the SSE stream exactly like the frontend does."""
    with client.stream("POST", "/api/extract", headers=headers, files={"file": (name, data, mime)}) as r:
        if r.status_code != 200:
            r.read()
            return r.status_code, r.json(), []
        body = "".join(r.iter_text())
    events = []
    for chunk in body.split("\n\n"):
        line = next((l for l in chunk.split("\n") if l.startswith("data:")), None)
        if line:
            events.append(json.loads(line[5:].strip()))
    return 200, None, events


# ---------- unit tests ----------


@pytest.mark.parametrize("raw,expected", [
    ("$1,234.50", 1234.5), ("1.234,50", 1234.5), ("12,50", 12.5), ("1,234", 1234.0),
    ("(45.00)", -45.0), ("-3", -3.0), (12, 12.0), (None, None), ("n/a", None), ("", None),
    ("USD 99", 99.0), ("1.234.567", 1234567.0), (True, None),
])
def test_to_number(raw, expected):
    assert to_number(raw) == expected


def test_invoice_coercion():
    inv = ExtractedInvoice.model_validate({
        "vendor_name": " Acme ", "invoice_number": 1042.0, "date": "null", "currency": "$",
        "total": "N1,200.50", "line_items": [{"description": "x", "total": "30", "quantity": 3}, "junk", None],
    })
    assert inv.vendor_name == "Acme" and inv.invoice_number == "1042" and inv.date is None
    assert inv.currency == "USD" and inv.total == 1200.5
    assert len(inv.line_items) == 1 and inv.line_items[0].unit_price == 10.0


def test_parse_json_variants():
    assert llm._parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert llm._parse_json('Sure! Here it is: {"a": 1} hope that helps') == {"a": 1}
    assert llm._parse_json('{"invoice": {"vendor_name": "X"}}') == {"vendor_name": "X"}
    with pytest.raises(ValueError):
        llm._parse_json("no json here")


def test_llm_retries_then_succeeds(monkeypatch):
    outputs = iter(["total garbage", GOOD_JSON])
    monkeypatch.setattr(llm, "_call_model", lambda text: next(outputs))
    assert llm.extract_invoice("some text").total == 161.25


def test_llm_gives_up_after_two_bad_outputs(monkeypatch):
    monkeypatch.setattr(llm, "_call_model", lambda text: "nope")
    with pytest.raises(llm.ExtractionError):
        llm.extract_invoice("some text")


def test_llm_rate_limit_is_friendly(monkeypatch):
    request = httpx.Request("POST", "https://api.groq.com/x")
    err = groq.RateLimitError("slow down", response=httpx.Response(429, request=request), body=None)

    def boom(text):
        raise err

    monkeypatch.setattr(llm, "_call_model", boom)
    with pytest.raises(llm.ExtractionError, match="busy"):
        llm.extract_invoice("some text")


def test_llm_empty_result_is_an_error(monkeypatch):
    monkeypatch.setattr(llm, "_call_model", lambda text: "{}")
    with pytest.raises(llm.ExtractionError, match="Couldn't find"):
        llm.extract_invoice("some text")


# ---------- document reading ----------


def test_text_pdf_uses_text_layer_not_ocr():
    result = ocr.extract_text(text_pdf(), "pdf")
    assert "INV-2026-0042" in result.text and result.ocr_pages == 0


def test_scanned_pdf_falls_back_to_ocr():
    result = ocr.extract_text(scanned_pdf(), "pdf")
    assert result.ocr_pages == 1
    assert "ACME" in result.text.upper() and "0042" in result.text


def test_png_goes_through_ocr():
    result = ocr.extract_text(invoice_png(), "image")
    assert "INVOICE" in result.text.upper() and "161.25" in result.text


def test_oversized_image_is_rejected_cleanly():
    big = Image.new("L", (9000, 9000), 255)  # 81 megapixels
    buf = io.BytesIO()
    big.save(buf, "PNG")
    with pytest.raises(ocr.DocumentError, match="could not be read"):
        ocr.extract_text(buf.getvalue(), "image")


def test_large_jpeg_is_handled():
    img = Image.new("RGB", (3000, 2000), "white")
    ImageDraw.Draw(img).text((100, 100), "INVOICE TOTAL 99.00", fill="black", font=ImageFont.load_default(size=90))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    assert "INVOICE" in ocr.extract_text(buf.getvalue(), "image").text.upper()


def test_multi_page_pdf_is_capped(monkeypatch):
    monkeypatch.setattr(ocr.settings, "max_pdf_pages", 3)
    result = ocr.extract_text(text_pdf(pages=5), "pdf")
    assert result.pages == 3 and result.truncated


def test_blank_pdf_is_rejected():
    doc = fitz.open()
    doc.new_page()
    with pytest.raises(ocr.DocumentError, match="No readable text"):
        ocr.extract_text(doc.tobytes(), "pdf")


def test_corrupt_pdf_is_rejected():
    with pytest.raises(ocr.DocumentError):
        ocr.extract_text(b"%PDF-1.4\nthis is not really a pdf", "pdf")


def test_sniff_kind():
    assert ocr.sniff_kind(text_pdf()) == "pdf"
    assert ocr.sniff_kind(invoice_png()) == "image"
    assert ocr.sniff_kind(b"MZ\x90\x00 an exe pretending") is None


# ---------- API flow ----------


def test_requires_auth(client):
    assert client.post("/api/extract", files={"file": ("a.pdf", text_pdf())}).status_code == 401
    assert client.get("/api/history").status_code == 401


def test_full_flow_text_pdf(client, headers, fake_llm):
    status, _, events = run_extract(client, headers, "invoice.pdf", text_pdf(), "application/pdf")
    assert status == 200
    stages = [e["stage"] for e in events]
    assert stages[0] == "uploading" and stages[-1] == "complete" and "error" not in stages
    assert [e["progress"] for e in events] == sorted(e["progress"] for e in events)

    record = events[-1]["result"]
    assert record["status"] == "complete" and record["filename"] == "invoice.pdf"
    data = record["extracted_data"]
    assert data["vendor_name"] == "Acme Supplies Ltd" and data["total"] == 161.25
    assert data["line_items"][1] == {"description": "Widget B", "quantity": 1.0, "unit_price": 100.0, "total": 100.0}
    assert "INV-2026-0042" in fake_llm[0]  # the PDF's text really reached the LLM step
    assert record["created_at"].endswith("+00:00") or record["created_at"].endswith("Z")

    history = client.get("/api/history", headers=headers).json()
    assert any(h["id"] == record["id"] for h in history)


def test_full_flow_image_and_scan(client, headers, fake_llm):
    for name, data in (("photo.png", invoice_png()), ("scan.pdf", scanned_pdf())):
        status, _, events = run_extract(client, headers, name, data)
        assert status == 200 and events[-1]["stage"] == "complete", (name, events)
    assert all("0042" in text for text in fake_llm)


def test_edit_then_export(client, headers, fake_llm):
    _, _, events = run_extract(client, headers, "inv.pdf", text_pdf())
    record = events[-1]["result"]
    edited = {**record["extracted_data"], "vendor_name": "=HYPERLINK(\"http://evil\")", "total": 999.99}

    r = client.patch(f"/api/extractions/{record['id']}", headers=headers, json=edited)
    assert r.status_code == 200 and r.json()["extracted_data"]["total"] == 999.99
    assert any(h["extracted_data"]["total"] == 999.99 for h in client.get("/api/history", headers=headers).json())

    csv_resp = client.get(f"/api/exports/csv?extraction_id={record['id']}", headers=headers)
    assert csv_resp.status_code == 200 and csv_resp.headers["content-type"].startswith("text/csv")
    text = csv_resp.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("vendor_name,invoice_number")
    assert len(text.strip().splitlines()) == 3  # header + 2 line items
    assert "999.99" in text
    assert "'=HYPERLINK" in text and ",=HYPERLINK" not in text and '"=HYPERLINK' not in text

    xl = client.get(f"/api/exports/excel?extraction_id={record['id']}", headers=headers)
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames == ["Invoice", "Line items"]
    vendor_cell = wb["Invoice"]["B2"]
    assert vendor_cell.data_type == "s" and vendor_cell.value.startswith("=HYPERLINK")  # stored as text, not a formula
    assert wb["Invoice"]["B8"].value == 999.99
    assert wb["Line items"].max_row == 3


def test_other_users_cannot_touch_my_records(client, headers, fake_llm):
    _, _, events = run_extract(client, headers, "mine.pdf", text_pdf())
    rid = events[-1]["result"]["id"]
    other = signup(client, "intruder@example.com")
    assert client.get(f"/api/exports/csv?extraction_id={rid}", headers=other).status_code == 404
    assert client.get(f"/api/exports/excel?extraction_id={rid}", headers=other).status_code == 404
    assert client.patch(f"/api/extractions/{rid}", headers=other, json={}).status_code == 404
    assert all(h["id"] != rid for h in client.get("/api/history", headers=other).json())


def test_upload_validation(client, headers, fake_llm):
    status, body, _ = run_extract(client, headers, "notes.txt", b"hello")
    assert status == 415 and "PDF, PNG or JPG" in body["detail"]

    status, body, _ = run_extract(client, headers, "fake.pdf", b"MZ\x90 definitely not a pdf")
    assert status == 415

    status, body, _ = run_extract(client, headers, "empty.pdf", b"")
    assert status == 400


def test_upload_too_large(client, headers, fake_llm, monkeypatch):
    from app.api.routes import extract as route
    monkeypatch.setattr(route.settings, "max_upload_mb", 0)
    status, body, _ = run_extract(client, headers, "big.pdf", text_pdf())
    assert status == 413 and "too large" in body["detail"]


def test_document_error_streams_error_and_marks_failed(client, headers, fake_llm):
    doc = fitz.open()
    doc.new_page()
    status, _, events = run_extract(client, headers, "blank.pdf", doc.tobytes())
    assert status == 200 and events[-1]["stage"] == "error"
    assert "No readable text" in events[-1]["message"]
    assert client.get("/api/history", headers=headers).json()[0]["status"] == "failed"


def test_llm_failure_streams_error(client, headers, monkeypatch):
    monkeypatch.setattr(llm, "_call_model", lambda text: "garbage")
    status, _, events = run_extract(client, headers, "inv.pdf", text_pdf())
    assert events[-1]["stage"] == "error" and "unreadable" in events[-1]["message"]

# ---------- stored documents: preview, delete, storage ----------


def invoice_jpg() -> bytes:
    buf = io.BytesIO()
    Image.open(io.BytesIO(invoice_png())).convert("RGB").save(buf, "JPEG", quality=90)
    return buf.getvalue()


def upload_ok(client, headers, name, data):
    status, body, events = run_extract(client, headers, name, data)
    assert status == 200 and events[-1]["stage"] == "complete", (status, body, events[-1:])
    return events[-1]["result"]


def history_ids(client, headers):
    return [h["id"] for h in client.get("/api/history", headers=headers).json()]


def test_original_is_stored_and_downloadable_for_pdf_png_and_jpg(client, fake_llm):
    h = signup(client, "storage-types@example.com")
    for name, data, mime in (
        ("a.pdf", text_pdf(), "application/pdf"),
        ("b.png", invoice_png(), "image/png"),
        ("c.jpg", invoice_jpg(), "image/jpeg"),
    ):
        record = upload_ok(client, h, name, data)
        assert record["has_file"] is True and record["file_size"] == len(data) and record["content_type"] == mime
        resp = client.get(f"/api/extractions/{record['id']}/file", headers=h)
        assert resp.status_code == 200 and resp.content == data
        assert resp.headers["content-type"].startswith(mime)
        assert resp.headers["x-content-type-options"] == "nosniff"
        assert name in resp.headers["content-disposition"]


def test_unicode_filename_download_header(client, fake_llm):
    h = signup(client, "unicode@example.com")
    record = upload_ok(client, h, "Facture été 2026.pdf", text_pdf())
    resp = client.get(f"/api/extractions/{record['id']}/file", headers=h)
    assert resp.status_code == 200 and "%C3%A9t%C3%A9" in resp.headers["content-disposition"]


def test_history_does_not_load_file_bytes(client, fake_llm):
    from sqlalchemy import event
    from app.db.base import engine

    h = signup(client, "efficient@example.com")
    upload_ok(client, h, "a.pdf", text_pdf())
    statements = []
    listener = lambda conn, cur, stmt, *a: statements.append(stmt)
    event.listen(engine, "before_cursor_execute", listener)
    try:
        assert client.get("/api/history", headers=h).status_code == 200
    finally:
        event.remove(engine, "before_cursor_execute", listener)
    selects = [s for s in statements if s.lstrip().upper().startswith("SELECT")]
    assert selects and not any("extraction_files.data" in s for s in selects)


def test_failed_extraction_stores_no_file(client, monkeypatch):
    monkeypatch.setattr(llm, "_call_model", lambda text: "garbage")
    h = signup(client, "nofile@example.com")
    _, _, events = run_extract(client, h, "x.pdf", text_pdf())
    assert events[-1]["stage"] == "error"
    record = client.get("/api/history", headers=h).json()[0]
    assert record["status"] == "failed" and record["has_file"] is False
    assert client.get(f"/api/extractions/{record['id']}/file", headers=h).status_code == 404


def test_delete_removes_everything(client, fake_llm):
    from app.db.base import SessionLocal
    from app.db.models import ExtractionFile

    h = signup(client, "deleter@example.com")
    record = upload_ok(client, h, "gone.pdf", text_pdf())
    rid = record["id"]

    assert client.delete(f"/api/extractions/{rid}", headers=h).status_code == 204
    assert rid not in history_ids(client, h)
    assert client.get(f"/api/extractions/{rid}/file", headers=h).status_code == 404
    assert client.get(f"/api/exports/csv?extraction_id={rid}", headers=h).status_code == 404
    assert client.delete(f"/api/extractions/{rid}", headers=h).status_code == 404  # already gone
    with SessionLocal() as db:
        assert db.get(ExtractionFile, rid) is None  # the stored bytes were deleted too


def test_cannot_delete_or_download_someone_elses_document(client, fake_llm):
    owner = signup(client, "owner2@example.com")
    other = signup(client, "other2@example.com")
    rid = upload_ok(client, owner, "private.pdf", text_pdf())["id"]
    assert client.delete(f"/api/extractions/{rid}", headers=other).status_code == 404
    assert client.get(f"/api/extractions/{rid}/file", headers=other).status_code == 404
    assert rid in history_ids(client, owner)  # untouched
    assert client.delete(f"/api/extractions/{rid}").status_code == 401  # and auth is required


def test_clear_failed_only_removes_failed(client, fake_llm, monkeypatch):
    h = signup(client, "clearer@example.com")
    good = upload_ok(client, h, "good.pdf", text_pdf())
    blank = fitz.open()
    blank.new_page()
    for _ in range(2):
        run_extract(client, h, "blank.pdf", blank.tobytes())
    assert sum(1 for r in client.get("/api/history", headers=h).json() if r["status"] == "failed") == 2

    resp = client.delete("/api/extractions/failed", headers=h)
    assert resp.status_code == 200 and resp.json() == {"deleted": 2}
    remaining = client.get("/api/history", headers=h).json()
    assert [r["id"] for r in remaining] == [good["id"]]
    assert client.delete("/api/extractions/failed", headers=h).json() == {"deleted": 0}


def test_storage_usage_and_quota(client, fake_llm, monkeypatch):
    from app.api.routes import extract as route

    h = signup(client, "quota@example.com")
    assert client.get("/api/storage", headers=h).json()["used_bytes"] == 0

    pdf = text_pdf()
    upload_ok(client, h, "one.pdf", pdf)
    usage = client.get("/api/storage", headers=h).json()
    assert usage["used_bytes"] == len(pdf) and usage["document_count"] == 1
    assert usage["limit_bytes"] == route.settings.max_storage_mb_per_user * 1024 * 1024

    # Shrink the quota so the next upload no longer fits: it is refused up front and nothing is recorded.
    monkeypatch.setattr(route.settings, "max_storage_mb_per_user", 0)
    before = history_ids(client, h)
    status, body, _ = run_extract(client, h, "two.pdf", pdf)
    assert status == 409 and "Storage limit reached" in body["detail"]
    assert history_ids(client, h) == before

    # Deleting frees the space again.
    monkeypatch.setattr(route.settings, "max_storage_mb_per_user", 25)
    assert client.delete(f"/api/extractions/{before[0]}", headers=h).status_code == 204
    assert client.get("/api/storage", headers=h).json()["used_bytes"] == 0


def test_storage_is_per_user(client, fake_llm):
    a = signup(client, "usage-a@example.com")
    b = signup(client, "usage-b@example.com")
    upload_ok(client, a, "a.pdf", text_pdf())
    assert client.get("/api/storage", headers=b).json() == {
        "used_bytes": 0, "limit_bytes": 25 * 1024 * 1024, "document_count": 0,
    }


def test_deleting_while_processing_is_handled(client, fake_llm):
    from app.api.routes import extract as route

    h = signup(client, "race@example.com")
    uid = client.get("/auth/me", headers=h).json()["id"]
    rid = route._create_record(uid, "racing.pdf")  # a record that is "processing"
    assert client.delete(f"/api/extractions/{rid}", headers=h).status_code == 204
    with pytest.raises(ocr.DocumentError, match="deleted"):
        route._complete_record(rid, ExtractedInvoice(vendor_name="X"), b"%PDF-1.4", "application/pdf")
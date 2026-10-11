"""Build CSV and Excel files from an extracted invoice."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.schemas.extraction import ExtractedInvoice

CSV_HEADERS = [
    "vendor_name", "invoice_number", "date", "currency",
    "description", "quantity", "unit_price", "line_total",
    "subtotal", "tax", "total",
]

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _csv_safe(value):
    """Stop spreadsheet apps from running text like '=HYPERLINK(...)' as a formula."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _blank(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)  # 2.0 -> 2
    return value


def to_csv(invoice: ExtractedInvoice) -> bytes:
    """One row per line item, with the invoice-level fields repeated on every row."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(CSV_HEADERS)

    def header_cells():
        return [invoice.vendor_name, invoice.invoice_number, invoice.date, invoice.currency]

    def footer_cells():
        return [invoice.subtotal, invoice.tax, invoice.total]

    rows = invoice.line_items or [None]
    for item in rows:
        line = [None, None, None, None] if item is None else [item.description, item.quantity, item.unit_price, item.total]
        cells = header_cells() + line + footer_cells()
        writer.writerow([_csv_safe(_blank(c)) for c in cells])

    # The BOM makes Excel open the file as UTF-8 (so names like "Café" survive).
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def _set_text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # force plain text so '=...' is never evaluated as a formula


def to_xlsx(invoice: ExtractedInvoice) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)

    summary = wb.active
    summary.title = "Invoice"
    summary.append(["Field", "Value"])
    fields = [
        ("Vendor name", invoice.vendor_name),
        ("Invoice number", invoice.invoice_number),
        ("Date", invoice.date),
        ("Currency", invoice.currency),
        ("Subtotal", invoice.subtotal),
        ("Tax", invoice.tax),
        ("Total", invoice.total),
    ]
    for label, value in fields:
        row = summary.max_row + 1
        summary.cell(row=row, column=1, value=label)
        _set_text(summary.cell(row=row, column=2), value)

    items = wb.create_sheet("Line items")
    items.append(["Description", "Quantity", "Unit price", "Total"])
    for item in invoice.line_items:
        row = items.max_row + 1
        _set_text(items.cell(row=row, column=1), item.description)
        items.cell(row=row, column=2, value=item.quantity)
        items.cell(row=row, column=3, value=item.unit_price)
        items.cell(row=row, column=4, value=item.total)

    for sheet, widths in ((summary, (20, 40)), (items, (50, 12, 14, 14))):
        for cell in sheet[1]:
            cell.font = bold
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()